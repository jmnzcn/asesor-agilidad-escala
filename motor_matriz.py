# -*- coding: utf-8 -*-
"""Motor de cálculo de la matriz de decisión (Nexus, LeSS, SAFe, Disciplined Agile).

Lee los criterios, pesos, niveles, puntajes y parámetros DIRECTAMENTE del archivo
Matriz_Decision_Agilidad_Escala.xlsx, de modo que el agente y la hoja de cálculo
usan siempre las mismas reglas. No contiene resultados de ningún caso.

Uso desde el agente (Code Interpreter):

    import sys; sys.path.insert(0, "/mnt/data")
    from motor_matriz import evaluar_caso
    r = evaluar_caso(ficha, titulo="Caso X", descripcion="...")
    print(r["reporte"])

`ficha` es un diccionario con una entrada por criterio:

    {"C1": {"valor": 22, "estado": "Explícito", "evidencia": "«22 equipos»", "supuesto": ""},
     "C2": {"valor": "Varias soluciones interdependientes", "estado": "Explícito", ...},
     ...
     "C6": {"valor": None, "estado": "Faltante", "evidencia": "sin evidencia", "supuesto": "..."}}
"""
from __future__ import annotations

import glob
import os
import unicodedata

from openpyxl import load_workbook

MARCOS = ["Nexus", "LeSS", "SAFe", "Disciplined Agile"]
CORTO = {"Nexus": "Nexus", "LeSS": "LeSS", "SAFe": "SAFe", "Disciplined Agile": "DA"}
# Regla R6 (c): parsimonia, mayor = más liviano de adoptar.
PARSIMONIA = {"Nexus": 4, "LeSS": 3, "Disciplined Agile": 2, "SAFe": 1}
ESTADOS = {"explicito": "Explícito", "inferido": "Inferido", "faltante": "Faltante"}
BANDAS_C1 = [(3, "1 a 3 equipos"), (9, "4 a 9 equipos"), (15, "10 a 15 equipos"),
             (50, "16 a 50 equipos"), (None, "Más de 50 equipos")]

FILA_CRITERIO_0 = 5    # 1_Criterios_Pesos: primera fila de criterios
FILA_REGLA_0 = 15      # 2_Escalas_Reglas: primera fila de reglas
FILA_PLANTILLA_0 = 6   # Plantilla: primera fila de captura


def _norm(texto) -> str:
    texto = "" if texto is None else str(texto)
    texto = unicodedata.normalize("NFD", texto.strip().lower())
    return "".join(c for c in texto if unicodedata.category(c) != "Mn")


def _fmt(x: float) -> str:
    """Dos decimales con redondeo «mitad hacia arriba», igual que Excel, y coma decimal."""
    from decimal import Decimal, ROUND_HALF_UP
    return str(Decimal(str(round(x, 6))).quantize(Decimal("0.01"), ROUND_HALF_UP)).replace(".", ",")


def _num(x: float) -> str:
    return f"{x:g}".replace(".", ",")


def buscar_excel(ruta: str | None = None) -> str:
    if ruta and os.path.exists(ruta):
        return ruta
    for patron in ("/mnt/data/*Matriz_Decision_Agilidad_Escala*.xlsx", "/mnt/data/*Matriz*.xlsx",
                   os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Matriz_Decision_Agilidad_Escala.xlsx"),
                   "Matriz_Decision_Agilidad_Escala.xlsx"):
        encontrados = sorted(glob.glob(patron))
        if encontrados:
            return encontrados[0]
    raise FileNotFoundError("No se encontró Matriz_Decision_Agilidad_Escala.xlsx. Súbalo a la conversación.")


# --------------------------------------------------------------------------- reglas
def cargar_reglas(ruta_excel: str | None = None) -> dict:
    """Lee pesos, niveles, puntajes y parámetros desde el Excel."""
    ruta = buscar_excel(ruta_excel)
    wb = load_workbook(ruta, data_only=True)
    wc, wr, wd = wb["1_Criterios_Pesos"], wb["2_Escalas_Reglas"], wb["3_Reglas_Decision"]

    criterios = {}
    fila = FILA_CRITERIO_0
    while wc.cell(fila, 1).value and str(wc.cell(fila, 1).value).startswith("C"):
        cid = wc.cell(fila, 1).value
        criterios[cid] = {"nombre": wc.cell(fila, 2).value, "peso": float(wc.cell(fila, 4).value),
                          "excluyente": _norm(wc.cell(fila, 5).value) == "si", "niveles": {}}
        fila += 1
    if abs(sum(c["peso"] for c in criterios.values()) - 1) > 1e-9:
        raise ValueError("Los pesos del Excel no suman 100 %.")

    fila = FILA_REGLA_0
    while wr.cell(fila, 2).value:
        cid, nivel = wr.cell(fila, 2).value, wr.cell(fila, 3).value
        criterios[cid]["niveles"][nivel] = {m: float(wr.cell(fila, 5 + i).value) for i, m in enumerate(MARCOS)}
        fila += 1

    return {"ruta": ruta, "criterios": criterios,
            "umbral_empate": float(wd["B4"].value), "neutro": float(wd["B5"].value)}


def banda_equipos(n: int) -> str:
    for tope, nombre in BANDAS_C1:
        if tope is None or n <= tope:
            return nombre
    raise ValueError(n)


# --------------------------------------------------------------------------- validación
def normalizar_ficha(ficha: dict, reglas: dict) -> dict:
    """Valida la ficha y resuelve el nivel aplicado de cada criterio."""
    errores, limpia = [], {}
    for cid, crit in reglas["criterios"].items():
        dato = dict(ficha.get(cid) or {})
        estado = ESTADOS.get(_norm(dato.get("estado")))
        valor = dato.get("valor")
        if estado is None:
            estado = "Faltante" if valor in (None, "") else None
        if estado is None:
            errores.append(f"{cid}: estado inválido «{dato.get('estado')}» (use Explícito, Inferido o Faltante).")
            continue
        nivel = None
        if estado != "Faltante":
            if cid == "C1":
                try:
                    valor = int(float(valor))
                    if valor < 1:
                        raise ValueError
                    nivel = banda_equipos(valor)
                except (TypeError, ValueError):
                    errores.append(f"C1: el valor debe ser un número entero de equipos ≥ 1 (recibido «{valor}»).")
                    continue
            else:
                opciones = {_norm(n): n for n in crit["niveles"]}
                nivel = opciones.get(_norm(valor))
                if nivel is None:
                    errores.append(f"{cid}: nivel «{valor}» no existe. Opciones: {', '.join(crit['niveles'])}.")
                    continue
                valor = nivel
        limpia[cid] = {"valor": valor if estado != "Faltante" else None, "estado": estado, "nivel": nivel,
                       "evidencia": dato.get("evidencia") or "sin evidencia", "supuesto": dato.get("supuesto") or ""}
    if errores:
        raise ValueError("Ficha inválida:\n- " + "\n- ".join(errores))
    return limpia


# --------------------------------------------------------------------------- cálculo
def calcular(ficha: dict, reglas: dict, exclusiones: dict | None = None) -> dict:
    """Aplica las reglas R1-R9 y devuelve el resultado completo y trazable."""
    exclusiones = {m: motivo for m, motivo in (exclusiones or {}).items() if motivo}
    crits = reglas["criterios"]
    filas, totales = [], {m: 0.0 for m in MARCOS}
    auto_excl = {m: [] for m in MARCOS}

    for cid, crit in crits.items():
        d = ficha[cid]
        if d["estado"] == "Faltante":
            puntos = {m: reglas["neutro"] for m in MARCOS}
        else:
            puntos = crit["niveles"][d["nivel"]]
            if crit["excluyente"]:
                for m in MARCOS:
                    if puntos[m] == 0:
                        auto_excl[m].append(cid)
        aportes = {m: crit["peso"] * puntos[m] for m in MARCOS}
        for m in MARCOS:
            totales[m] += aportes[m]
        filas.append({"id": cid, "criterio": crit["nombre"], "peso": crit["peso"], "excluyente": crit["excluyente"],
                      "nivel": d["nivel"] or "(faltante)", "estado": d["estado"], "valor": d["valor"],
                      "evidencia": d["evidencia"], "supuesto": d["supuesto"], "puntos": puntos, "aportes": aportes})

    totales = {m: round(v, 6) for m, v in totales.items()}
    excluido = {m: bool(auto_excl[m]) or m in exclusiones for m in MARCOS}
    elegibles = [m for m in MARCOS if not excluido[m]]

    resultado = {"filas": filas, "totales": totales, "excluido": excluido, "auto_excl": auto_excl,
                 "excl_manual": exclusiones, "umbral": reglas["umbral_empate"], "reglas_excel": reglas["ruta"]}
    pesos = {"explícito": 0.0, "inferido": 0.0, "faltante": 0.0}
    for f in filas:
        pesos[_norm(f["estado"]).replace("explicito", "explícito")] += f["peso"]
    resultado["cobertura"] = pesos

    if not elegibles:  # R9
        resultado.update(recomendacion=None, segundo=None, brecha=None, empate=False, en_empate=[],
                         robustez="Sin recomendación", peor_brecha={}, aportes_vs_segundo=[], riesgos=[])
        return resultado

    # R5 y R6
    maximo = max(totales[m] for m in elegibles)
    en_empate = [m for m in elegibles if round(totales[m], 4) >= round(maximo - reglas["umbral_empate"], 4)]
    riesgo = {m: sum(1 for f in filas if f["puntos"][m] <= 1) for m in MARCOS}
    estructural = {m: sum(f["aportes"][m] for f in filas if f["excluyente"]) for m in MARCOS}
    clave = {m: 10000 - riesgo[m] * 1000 + estructural[m] * 100 + PARSIMONIA[m] for m in en_empate}
    rec = max(en_empate, key=lambda m: clave[m])
    orden = sorted(elegibles, key=lambda m: -totales[m])
    segundo = next((m for m in orden if m != rec), None)

    # R7: brecha en el peor caso por cada competidor elegible
    peor = {}
    faltantes = [f for f in filas if f["estado"] == "Faltante"]
    for x in elegibles:
        if x == rec:
            continue
        ajuste = 0.0
        for f in faltantes:
            niveles = crits[f["id"]]["niveles"].values()
            ajuste += f["peso"] * min(p[rec] - p[x] for p in niveles)
        peor[x] = {"brecha": totales[rec] - totales[x], "ajuste": ajuste,
                   "peor": round(totales[rec] - totales[x] + ajuste, 6)}
    if not faltantes:
        robustez = "Robusta: no hay datos faltantes"
    elif any(v["peor"] <= 0 for v in peor.values()):
        robustez = "SENSIBLE: los datos faltantes pueden cambiar la recomendación; se requiere aclaración"
    else:
        robustez = "Robusta: ningún valor posible de los datos faltantes cambia la recomendación"

    aportes_vs = []
    if segundo:
        for f in filas:
            dif = f["aportes"][rec] - f["aportes"][segundo]
            if abs(dif) > 1e-9:
                aportes_vs.append((f["id"], f["criterio"], round(dif, 4)))
        aportes_vs.sort(key=lambda t: -t[2])

    resultado.update(
        recomendacion=rec, segundo=segundo,
        brecha=round(totales[rec] - totales[segundo], 6) if segundo else None,
        empate=len(en_empate) > 1, en_empate=en_empate, riesgo=riesgo, estructural=estructural,
        robustez=robustez, peor_brecha=peor, aportes_vs_segundo=aportes_vs,
        riesgos=[(f["id"], f["criterio"], f["puntos"][rec]) for f in filas
                 if f["estado"] != "Faltante" and f["puntos"][rec] <= 2],
        faltantes=[f["id"] for f in faltantes])
    return resultado


# --------------------------------------------------------------------------- reporte
def reporte_markdown(r: dict) -> str:
    L = ["### 1. Ficha de captura", "",
         "| ID | Criterio | Peso | Evidencia citada | Nivel | Estado | Supuesto |", "|---|---|---|---|---|---|---|"]
    for f in r["filas"]:
        nivel = f["nivel"] if f["id"] != "C1" or f["valor"] is None else f"{f['valor']} equipos → {f['nivel']}"
        L.append(f"| {f['id']} | {f['criterio']} | {f['peso'] * 100:.0f} % | {f['evidencia']} | {nivel} | {f['estado']} | {f['supuesto'] or '—'} |")
    L += ["", "### 2. Cálculo (compatibilidad → peso × compatibilidad)", "",
          "| ID | Peso | " + " | ".join(CORTO[m] for m in MARCOS) + " |", "|---|---|" + "---|" * 4]
    for f in r["filas"]:
        celdas = [f"{_num(f['puntos'][m])} → {_fmt(f['aportes'][m])}" for m in MARCOS]
        L.append(f"| {f['id']} | {_fmt(f['peso'])} | " + " | ".join(celdas) + " |")
    L.append("| **Total** | **1,00** | " + " | ".join(f"**{_fmt(r['totales'][m])}**" for m in MARCOS) + " |")

    L += ["", "### 3. Restricciones críticas"]
    hubo = False
    for m in MARCOS:
        if r["auto_excl"][m]:
            L.append(f"- **{m}: excluido** (R3) por compatibilidad 0 en criterio excluyente {', '.join(r['auto_excl'][m])}.")
            hubo = True
        if m in r["excl_manual"]:
            L.append(f"- **{m}: excluido** (R4, exclusión manual): {r['excl_manual'][m]}.")
            hubo = True
    if not hubo:
        L.append("- Ninguno: los cuatro marcos son elegibles.")

    L += ["", "### 4. Recomendación"]
    if r["recomendacion"] is None:
        L.append("**Ninguno compatible** (R9): los cuatro marcos quedaron excluidos; revise el caso.")
        return "\n".join(L)
    L.append(f"- **Marco recomendado: {r['recomendacion']} ({_fmt(r['totales'][r['recomendacion']])} / 5,00)**")
    if r["segundo"]:
        L.append(f"- Segundo lugar elegible: {r['segundo']} ({_fmt(r['totales'][r['segundo']])}); brecha {_fmt(r['brecha'])}.")
    if r["empate"]:
        riesgos_txt = ", ".join("%s %d" % (m, r["riesgo"][m]) for m in r["en_empate"])
        estruct_txt = ", ".join("%s %s" % (m, _fmt(r["estructural"][m])) for m in r["en_empate"])
        L.append(f"- Empate técnico (≤ {_fmt(r['umbral'])}) entre {', '.join(r['en_empate'])}: se aplicó R6 "
                 f"(criterios ≤ 1: {riesgos_txt}; estructural: {estruct_txt}; luego parsimonia).")
    else:
        L.append(f"- Empate técnico: No (umbral {_fmt(r['umbral'])}).")

    L += ["", "### 5. Robustez ante datos faltantes (R7)", f"**{r['robustez']}**"]
    if r["faltantes"]:
        L.append(f"Criterios faltantes: {', '.join(r['faltantes'])} (puntaje neutro 2,5).")
        L += ["", "| Competidor | Brecha actual | Ajuste peor caso | Brecha en peor caso |", "|---|---|---|---|"]
        for x, v in r["peor_brecha"].items():
            marca = " ⚠️" if v["peor"] <= 0 else ""
            L.append(f"| {x} | {_fmt(v['brecha'])} | {_fmt(v['ajuste'])} | {_fmt(v['peor'])}{marca} |")

    L += ["", "### 6. Criterios que explican la diferencia con el segundo lugar"]
    for cid, nombre, dif in r["aportes_vs_segundo"][:4]:
        signo = "+" if dif > 0 else ""
        L.append(f"- {cid} {nombre}: {signo}{_fmt(dif)}")
    if r["riesgos"]:
        L.append("")
        L.append("Riesgos de adopción del marco recomendado (compatibilidad ≤ 2): " +
                 "; ".join(f"{cid} {n} ({_num(p)})" for cid, n, p in r["riesgos"]) + ".")
    c = r["cobertura"]
    L += ["", f"Cobertura de datos por peso: explícito {c['explícito']:.0%} · inferido {c['inferido']:.0%} · faltante {c['faltante']:.0%}.",
          f"Reglas leídas de: {os.path.basename(r['reglas_excel'])}."]
    return "\n".join(L)


# --------------------------------------------------------------------------- salidas
def grafico(r: dict, ruta_png: str, titulo: str = "") -> str:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    colores = {"Nexus": "#4C72B0", "LeSS": "#55A868", "SAFe": "#C44E52", "Disciplined Agile": "#8172B2"}
    fig, ax = plt.subplots(figsize=(8, 4.2), dpi=150)
    valores = [r["totales"][m] for m in MARCOS]
    barras = ax.bar([CORTO[m] for m in MARCOS], valores, color=[colores[m] for m in MARCOS],
                    edgecolor="black", linewidth=0.6)
    for b, m in zip(barras, MARCOS):
        if r["excluido"][m]:
            b.set_hatch("///")
            b.set_alpha(0.45)
        etiqueta = _fmt(r["totales"][m]) + ("\nexcluido" if r["excluido"][m] else "")
        if m == r.get("recomendacion"):
            etiqueta += "\n★ recomendado"
            b.set_linewidth(2.2)
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.08, etiqueta, ha="center", va="bottom", fontsize=9)
    ax.set_ylim(0, 5.9)
    ax.set_ylabel("Puntaje ponderado (máx. 5,00)")
    ax.set_title(titulo or "Resultado de la matriz de decisión", fontsize=11)
    ax.spines[["top", "right"]].set_visible(False)
    ax.text(0.99, 0.97, r.get("robustez", "").split(":")[0], transform=ax.transAxes, ha="right", va="top",
            fontsize=8, color="#7F0000" if r.get("robustez", "").startswith("SENSIBLE") else "#1F5F1F")
    fig.tight_layout()
    fig.savefig(ruta_png)
    plt.close(fig)
    return ruta_png


def exportar_excel(ficha: dict, reglas: dict, ruta_salida: str, titulo: str, descripcion: str,
                   exclusiones: dict | None = None) -> str:
    """Copia la matriz y llena la hoja Plantilla con la ficha para auditar en Excel."""
    wb = load_workbook(reglas["ruta"])
    ws = wb["Plantilla"]
    ws["A1"] = titulo
    ws["B2"] = descripcion
    for i, cid in enumerate(reglas["criterios"]):
        fila = FILA_PLANTILLA_0 + i
        d = ficha[cid]
        ws[f"E{fila}"] = d["evidencia"]
        ws[f"F{fila}"] = d["valor"]
        ws[f"G{fila}"] = d["estado"]
        ws[f"H{fila}"] = d["supuesto"]
    fila_res = FILA_PLANTILLA_0 + len(reglas["criterios"]) + 4
    for k, m in enumerate(MARCOS):
        motivo = (exclusiones or {}).get(m)
        ws[f"D{fila_res + k}"] = "Sí" if motivo else "No"
        ws[f"E{fila_res + k}"] = motivo or None
    ws.title = "Caso_Evaluado"
    wb.move_sheet("Caso_Evaluado", offset=-(wb.sheetnames.index("Caso_Evaluado") - 1))
    wb.active = 1
    for hoja in wb.worksheets:
        hoja.sheet_view.tabSelected = hoja.title == "Caso_Evaluado"
    wb.calculation.fullCalcOnLoad = True  # Excel recalcula todas las fórmulas al abrir
    wb.save(ruta_salida)
    return ruta_salida


def _nombre_archivo(titulo: str) -> str:
    base = "".join(c if c.isalnum() else "_" for c in _norm(titulo))[:40].strip("_")
    return base or "caso"


def evaluar_caso(ficha: dict, titulo: str = "Caso evaluado", descripcion: str = "",
                 exclusiones: dict | None = None, ruta_excel: str | None = None,
                 carpeta_salida: str | None = None) -> dict:
    """Punto de entrada del agente: valida, calcula, grafica y exporta."""
    reglas = cargar_reglas(ruta_excel)
    limpia = normalizar_ficha(ficha, reglas)
    r = calcular(limpia, reglas, exclusiones)
    carpeta = carpeta_salida or ("/mnt/data/salidas" if os.path.isdir("/mnt/data") else "salidas")
    os.makedirs(carpeta, exist_ok=True)
    base = _nombre_archivo(titulo)
    r["grafico"] = grafico(r, os.path.join(carpeta, f"{base}_grafico.png"), titulo)
    r["excel"] = exportar_excel(limpia, reglas, os.path.join(carpeta, f"{base}_matriz.xlsx"), titulo, descripcion, exclusiones)
    r["reporte"] = reporte_markdown(r)
    return r


if __name__ == "__main__":
    import json
    import sys

    if len(sys.argv) < 2:
        print("Uso: python motor_matriz.py ficha.json [ruta_excel]")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as fh:
        entrada = json.load(fh)
    res = evaluar_caso(entrada["ficha"], entrada.get("titulo", "Caso"), entrada.get("descripcion", ""),
                       entrada.get("exclusiones"), sys.argv[2] if len(sys.argv) > 2 else None)
    print(res["reporte"])
    print(f"\nGráfico: {res['grafico']}\nExcel: {res['excel']}")

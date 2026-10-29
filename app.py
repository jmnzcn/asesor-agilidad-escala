# -*- coding: utf-8 -*-
"""Asesor de Agilidad a Escala (PF-3872, UCR).

La IA (Gemini) lee el caso y llena la ficha de captura; motor_matriz.py calcula
con las reglas de Matriz_Decision_Agilidad_Escala.xlsx; el Excel descargable
permite auditar el resultado. Ejecutar con:  streamlit run app.py
"""
from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

import graficos
import motor_matriz as motor

RAIZ = Path(__file__).resolve().parent
EXCEL = str(RAIZ / "Matriz_Decision_Agilidad_Escala.xlsx")
# Enunciados de los casos del curso. Solo el texto: la ficha la llena la IA (o el evaluador) desde cero.
CASOS = {
    "Caso 1 · Banco regional": {"titulo": "Caso 1. Banco regional",
        "texto": "Banco en seis países; 22 equipos distribuidos; varios productos digitales conectados; dependencias con arquitectura, datos, seguridad y operaciones; regulación y necesidad de vincular iniciativas con estrategia e inversiones."},
    "Caso 2 · Aseguradora multinacional": {"titulo": "Caso 2. Aseguradora multinacional",
        "texto": "Modernización de pólizas, reclamaciones y riesgos; 25 equipos y aproximadamente 300 profesionales; varias soluciones interdependientes; coordinación de proveedores, negocio y tecnología; planificación transversal y visibilidad de portafolio."},
    "Caso 3 · Servicios profesionales": {"titulo": "Caso 3. Servicios profesionales",
        "texto": "Seis equipos para proyectos de clientes distintos; algunos iterativos, otros de mantenimiento y otros híbridos; exigencias contractuales variadas; autonomía y necesidad de adaptar prácticas sin imponer un único proceso."},
    "Prueba 4A · Cambio de condición": {"titulo": "Prueba 4A. Cambio de condición (variante del Caso 1)",
        "texto": "El mismo banco regional limita la iniciativa a su aplicación de banca móvil. Participan 7 equipos Scrum sobre un único producto con un solo Product Backlog y un Product Owner. Las dependencias se dan entre los equipos que integran el mismo incremento. La regulación bancaria se mantiene. El presupuesto ya está aprobado y no se requiere gestión de portafolio adicional. Todo el trabajo es desarrollo iterativo de nuevas funcionalidades."},
    "Prueba 4B · Datos faltantes": {"titulo": "Prueba 4B. Datos faltantes decisivos",
        "texto": "Una aseguradora quiere modernizar su plataforma de reclamaciones con varios equipos internos y un proveedor externo de desarrollo. Todavía no se ha definido cuántos equipos participarán ni si la iniciativa se gestionará dentro del portafolio corporativo. Los equipos internos trabajan con Scrum."},
}
ESTADOS = ["Explícito", "Inferido", "Faltante"]
COLORES = {"Nexus": "#3566c8", "LeSS": "#24875c", "SAFe": "#c4414f", "Disciplined Agile": "#7454b8"}
HORA_CR = timezone(timedelta(hours=-6))  # Costa Rica, sin horario de verano
TEMPERATURA = 0.1

st.set_page_config(page_title="Asesor de Agilidad a Escala", page_icon="🧭", layout="wide")


# --------------------------------------------------------------------------- datos
@st.cache_data
def reglas() -> dict:
    return motor.cargar_reglas(EXCEL)


def ficha_vacia() -> dict:
    return {cid: {"valor": None, "estado": "Faltante", "evidencia": "", "supuesto": ""} for cid in reglas()["criterios"]}


def cargar_en_estado(titulo: str, texto: str, ficha: dict, origen: str, exclusiones: dict | None = None):
    st.session_state.titulo = titulo
    st.session_state.texto = texto
    st.session_state.ficha = json.loads(json.dumps(ficha))
    st.session_state.excl = dict(exclusiones or {})
    st.session_state.origen = origen
    st.session_state.justificacion = ""
    st.session_state.ia = None      # respuesta cruda de la IA para esta ficha (transparencia)
    st.session_state.antes = None   # resultado previo a una aclaración (comparación antes → después)
    st.session_state.version = st.session_state.get("version", 0) + 1  # renueva los widgets


def ahora() -> str:
    return datetime.now(HORA_CR).strftime("%d/%m/%Y %H:%M:%S")


if "bitacora" not in st.session_state:
    st.session_state.bitacora = []
if "ficha" not in st.session_state:
    primero = next(iter(CASOS))
    c = CASOS[primero]
    cargar_en_estado(c["titulo"], c["texto"], ficha_vacia(), "Ficha vacía")
    st.session_state.caso_sel = primero


# --------------------------------------------------------------------------- IA (Gemini)
def _secreto(nombre: str) -> str:
    """Lee un secreto de Streamlit o, si no hay archivo de secretos, de una variable de entorno."""
    try:
        valor = st.secrets.get(nombre, "")
    except Exception:  # sin archivo de secretos
        valor = ""
    return valor or os.environ.get(nombre, "")


def gemini():
    from google import genai
    return genai.Client(api_key=_secreto("GEMINI_API_KEY")), _secreto("GEMINI_MODEL") or "gemini-2.5-flash"


def _secretos_ok() -> bool:
    return bool(_secreto("GEMINI_API_KEY"))


def reglas_para_prompt() -> str:
    L = []
    for cid, c in reglas()["criterios"].items():
        L.append(f"{cid} {c['nombre']} (peso {c['peso'] * 100:.0f} %{', excluyente' if c['excluyente'] else ''})")
        if cid == "C1":
            L.append("  valor = número entero de equipos (si solo hay personas: equipos ≈ personas / 10, estado Inferido).")
            L += [f"  bandas: {', '.join(c['niveles'])}"]
        else:
            L += [f'  - "{n}"' for n in c["niveles"]]
    return "\n".join(L)


def disparadores() -> dict:
    """Texto del disparador de cada nivel, leído del Excel."""
    wb = load_workbook(EXCEL, data_only=True)
    wr = wb["2_Escalas_Reglas"]
    out, fila = {}, motor.FILA_REGLA_0
    while wr.cell(fila, 2).value:
        out[(wr.cell(fila, 2).value, wr.cell(fila, 3).value)] = wr.cell(fila, 4).value
        fila += 1
    return out


DISP = disparadores()


def prompt_extraccion(texto: str) -> str:
    niveles = "\n".join(f'{cid} "{niv}": {d}' for (cid, niv), d in DISP.items())
    return ("Eres el analista de una matriz de decisión que recomienda Nexus, LeSS, SAFe o Disciplined Agile. "
            "Tu única tarea es LEER el caso y llenar la ficha de captura; no calcules ni recomiendes.\n\n"
            "CRITERIOS:\n" + reglas_para_prompt() + "\n\nDISPARADOR DE CADA NIVEL (copia el nombre exacto):\n" + niveles +
            "\n\nESTADOS: \"Explícito\" (el caso lo dice), \"Inferido\" (evidencia indirecta; escribe el supuesto), "
            "\"Faltante\" (no hay base; valor null). No inventes datos. Si el caso declara algo que impide un marco "
            "(p. ej. que no se usará Scrum), ponlo en exclusiones con la frase citada.\n\n"
            "Responde solo con JSON: {\"ficha\":{\"C1\":{\"valor\":22,\"estado\":\"Explícito\",\"evidencia\":\"«cita»\",\"supuesto\":\"\"}, "
            "... \"C7\":{...}}, \"exclusiones\":{\"<marco>\":\"«cita»\"}}\n\nCASO:\n" + texto[:6000])


def resumen(r: dict) -> str:
    L = ["Totales: " + "; ".join(f"{m} {motor._fmt(r['totales'][m])}{' (excluido)' if r['excluido'][m] else ''}" for m in motor.MARCOS),
         f"Recomendado: {r['recomendacion']}; segundo: {r['segundo']}; brecha {motor._fmt(r['brecha'] or 0)}"
         + ("; empate técnico resuelto por R6" if r["empate"] else ""),
         "Robustez: " + r["robustez"]]
    for m in motor.MARCOS:
        if r["auto_excl"][m]:
            L.append(f"{m} excluido por compatibilidad 0 en {', '.join(r['auto_excl'][m])}")
        if m in r["excl_manual"]:
            L.append(f"{m} excluido por restricción: {r['excl_manual'][m]}")
    L.append("Criterios que más separan: " + "; ".join(f"{c} {n} {motor._fmt(d)}" for c, n, d in r["aportes_vs_segundo"][:4]))
    if r["riesgos"]:
        L.append("Riesgos (compatibilidad ≤ 2): " + "; ".join(f"{c} {n}" for c, n, _ in r["riesgos"]))
    L.append("Ficha: " + " | ".join(f"{f['id']} {f['criterio']} = {f['nivel']} [{f['estado']}]" for f in r["filas"]))
    return "\n".join(L)


def prompt_justificacion(texto: str, r: dict) -> str:
    final = (f"4) La recomendación es PROVISIONAL: termina con 'Preguntas:' y 1 a 3 preguntas sobre {', '.join(r['faltantes'])}."
             if r.get("robustez", "").startswith("SENSIBLE") else
             "4) Termina con 'Supuestos:' listando los datos inferidos o faltantes, o 'ninguno'.")
    return ("Eres el asesor de agilidad a escala de un equipo de maestría. La matriz ya calculó el resultado; no cambies ningún número. "
            "Escribe en español, máximo 180 palabras, con: 1) por qué gana " + str(r["recomendacion"]) +
            " usando los 2-3 criterios que más lo separan del segundo; 2) por qué los demás quedan atrás o fueron excluidos; "
            "3) riesgos de adopción y una recomendación práctica. " + final +
            "\n\nRESULTADO:\n" + resumen(r) + "\n\nCASO:\n" + (texto or "")[:4000])


def mensaje_error(e: Exception) -> str:
    t = str(e)
    if "429" in t or "RESOURCE_EXHAUSTED" in t:
        return ("La IA recibió demasiadas consultas en este minuto (límite del plan gratuito de Gemini). "
                "Espere un minuto y vuelva a intentarlo, o llene la ficha a mano: el cálculo y el Excel funcionan igual.")
    return f"No se pudo completar la lectura con IA: {t}"


def analizar_con_ia(texto: str, titulo: str, origen: str = "Ficha generada por IA"):
    """Envía el caso a la IA, valida la ficha que devuelve y la carga; guarda la respuesta cruda."""
    cliente, modelo = gemini()
    from google.genai import types
    prompt = prompt_extraccion(texto)
    resp = cliente.models.generate_content(
        model=modelo, contents=prompt,
        config=types.GenerateContentConfig(response_mime_type="application/json", temperature=TEMPERATURA))
    salida = json.loads(resp.text)
    nueva, ajustes = ficha_vacia(), []
    for cid in nueva:
        d = (salida.get("ficha") or {}).get(cid) or {}
        nueva[cid] = {"valor": d.get("valor"), "estado": d.get("estado") or "Faltante",
                      "evidencia": str(d.get("evidencia") or ""), "supuesto": str(d.get("supuesto") or "")}
    # un nivel inválido pasa a Faltante y se informa
    for cid in nueva:
        try:
            motor.normalizar_ficha({**ficha_vacia(), cid: nueva[cid]}, reglas())
        except ValueError:
            ajustes.append(f"{cid}: la IA propuso «{nueva[cid]['valor']}», que no es un nivel de la matriz; quedó como Faltante.")
            nueva[cid].update(supuesto=f"La IA propuso «{nueva[cid]['valor']}», que no es un nivel válido.",
                              valor=None, estado="Faltante")
    excl = {}
    for k, val in (salida.get("exclusiones") or {}).items():
        m = next((x for x in motor.MARCOS if motor._norm(x) == motor._norm(k) or motor._norm(motor.CORTO[x]) == motor._norm(k)), None)
        if m and val:
            excl[m] = str(val)
    cargar_en_estado(titulo, texto, nueva, origen, excl)
    st.session_state.ia = {"modelo": modelo, "hora": ahora(), "temperatura": TEMPERATURA,
                           "prompt": prompt, "respuesta": resp.text, "ajustes": ajustes}
    st.session_state.pedir_justificacion = True


# --------------------------------------------------------------------------- bitácora
def firma(limpia: dict) -> str:
    return json.dumps([st.session_state.texto, limpia, st.session_state.excl], sort_keys=True, ensure_ascii=False, default=str)


def instantanea(r: dict) -> dict:
    """Lo necesario de un resultado para registrarlo y compararlo después."""
    return {"totales": dict(r["totales"]), "excluido": dict(r["excluido"]), "recomendacion": r["recomendacion"],
            "segundo": r["segundo"], "brecha": r["brecha"], "robustez": r["robustez"], "faltantes": list(r.get("faltantes") or []),
            "filas": [{k: f[k] for k in ("id", "criterio", "nivel", "estado", "valor", "evidencia", "supuesto", "aportes")}
                      for f in r["filas"]]}


def registrar(r: dict, limpia: dict):
    """Agrega la ejecución actual a la bitácora; si es la misma ficha que la última, solo actualiza la justificación."""
    b, f = st.session_state.bitacora, firma(limpia)
    if b and b[-1]["firma"] == f:
        b[-1]["justificacion"] = st.session_state.justificacion
        return
    ia = st.session_state.get("ia") or {}
    b.append({"n": len(b) + 1, "hora": ahora(), "caso": st.session_state.titulo, "texto": st.session_state.texto,
              "origen": st.session_state.origen, "modelo": ia.get("modelo", "manual (sin IA)"),
              "justificacion": st.session_state.justificacion, "exclusiones": dict(st.session_state.excl),
              "firma": f, **instantanea(r)})


def _txt_total(e: dict) -> str:
    rec = e["recomendacion"]
    return f"{rec} {motor._fmt(e['totales'][rec])}" if rec else "Ninguno (R9)"


def _robustez_corta(e: dict) -> str:
    return "Sensible" if e["robustez"].startswith("SENSIBLE") else "Robusta"


def _delta(d: float) -> str:
    return "—" if abs(d) < 0.005 else f"{'+' if d > 0 else ''}{motor._fmt(d)}"


def panel_bitacora():
    st.divider()
    st.header("Bitácora de ejecuciones")
    b = st.session_state.bitacora
    st.caption("Cada análisis con IA se registra aquí con su hora, modelo, puntajes, robustez y supuestos; también puede registrar "
               "una ficha llenada a mano. La bitácora vive mientras esta pestaña esté abierta: descárguela como evidencia.")
    if not b:
        st.info("Aún no hay ejecuciones. Pulse «Analizar con IA» en un caso.")
        return
    st.dataframe(pd.DataFrame([{
        "#": e["n"], "Hora": e["hora"], "Caso": e["caso"], "Origen": e["origen"], "Modelo": e["modelo"],
        **{motor.CORTO[m]: motor._fmt(e["totales"][m]) + (" (excl.)" if e["excluido"][m] else "") for m in motor.MARCOS},
        "Recomendación": e["recomendacion"] or "Ninguno (R9)", "Robustez": _robustez_corta(e),
        "Faltantes": ", ".join(e["faltantes"]) or "—"} for e in b]), hide_index=True, width="stretch")
    exportar = pd.DataFrame([{
        "n": e["n"], "hora": e["hora"], "caso": e["caso"], "descripcion": e["texto"], "origen": e["origen"], "modelo": e["modelo"],
        **{f"total_{motor.CORTO[m]}": round(e["totales"][m], 4) for m in motor.MARCOS},
        "recomendacion": e["recomendacion"] or "Ninguno (R9)", "robustez": e["robustez"], "faltantes": ", ".join(e["faltantes"]),
        "ficha": " | ".join(f"{f['id']} {f['nivel']} [{f['estado']}]" for f in e["filas"]),
        "supuestos": " | ".join(f"{f['id']}: {f['supuesto']}" for f in e["filas"] if f["supuesto"]),
        "justificacion": e["justificacion"]} for e in b])
    c1, c2, c3 = st.columns([1, 1, 1])
    c1.download_button("Descargar bitácora (CSV)", exportar.to_csv(index=False).encode("utf-8-sig"),
                       file_name="Bitacora_ejecuciones.csv", mime="text/csv", type="primary")
    c2.download_button("Descargar bitácora (JSON)", json.dumps([{k: v for k, v in e.items() if k != "firma"} for e in b],
                       ensure_ascii=False, indent=2, default=str), file_name="Bitacora_ejecuciones.json", mime="application/json")
    if c3.button("Vaciar bitácora"):
        st.session_state.bitacora = []
        for k in ("cmp_a", "cmp_b"):
            st.session_state.pop(k, None)
        st.rerun()

    if len(b) < 2:
        st.caption("Con dos o más ejecuciones podrá compararlas lado a lado.")
        return
    st.subheader("Comparar dos ejecuciones")
    st.caption("Por ejemplo, el Caso 1 frente a la prueba 4A (cambio de condición) o la 4B antes y después de responder al agente: "
               "se ve qué criterios cambiaron y cuánto movió cada uno el puntaje de cada marco.")
    etiquetas = [f"#{e['n']} · {e['caso']} · {_txt_total(e)}" for e in b]
    ca, cb = st.columns(2)
    for k in ("cmp_a", "cmp_b"):
        if st.session_state.get(k, 0) >= len(b):
            st.session_state.pop(k)
    ia_ = ca.selectbox("Ejecución A", range(len(b)), index=len(b) - 2, format_func=lambda i: etiquetas[i], key="cmp_a")
    ib_ = cb.selectbox("Ejecución B", range(len(b)), index=len(b) - 1, format_func=lambda i: etiquetas[i], key="cmp_b")
    A, B = b[ia_], b[ib_]
    cambios = [(fa, fb) for fa, fb in zip(A["filas"], B["filas"]) if (fa["nivel"], fa["estado"]) != (fb["nivel"], fb["estado"])]
    st.markdown(f"**A:** {_txt_total(A)} ({_robustez_corta(A)}) → **B:** {_txt_total(B)} ({_robustez_corta(B)}). "
                + (f"Cambiaron {len(cambios)} de 7 criterios." if cambios else "La ficha es la misma en ambas ejecuciones."))
    st.plotly_chart(graficos.comparar(A, B, f"A · #{A['n']}", f"B · #{B['n']}"), width="stretch",
                    config={"displayModeBar": False, "locale": "es"})
    filas = []
    for fa, fb in zip(A["filas"], B["filas"]):
        filas.append({"ID": fa["id"], "Criterio": fa["criterio"], "A": f"{fa['nivel']} ({fa['estado']})",
                      "B": f"{fb['nivel']} ({fb['estado']})", "Cambió": "●" if (fa, fb) in cambios else "",
                      **{f"Δ {motor.CORTO[m]}": _delta(fb["aportes"][m] - fa["aportes"][m]) for m in motor.MARCOS}})
    filas.append({"ID": "Total", "Criterio": "", "A": "", "B": "", "Cambió": "",
                  **{f"Δ {motor.CORTO[m]}": _delta(B["totales"][m] - A["totales"][m]) for m in motor.MARCOS}})
    st.dataframe(pd.DataFrame(filas), hide_index=True, width="stretch")
    st.caption("Δ = aporte en B − aporte en A (peso × compatibilidad). La suma de cada columna explica el cambio del total de ese marco.")


# --------------------------------------------------------------------------- interfaz
st.markdown(
    """
    <style>
      .block-container {padding-top: 1.6rem;}
      .eyebrow {font-family: ui-monospace, Consolas, monospace; font-size: 12px; letter-spacing: .06em;
                text-transform: uppercase; opacity: .7;}
      .veredicto {padding: 14px 18px; border-radius: 8px; border-left: 6px solid var(--c);
                  background: rgba(127,127,127,.08);}
      .veredicto .marco {font-size: 30px; font-weight: 800; line-height: 1.1;}
      .pill {display:inline-block; padding: 2px 10px; border-radius: 999px; font-size: 12px; font-weight: 700; margin-right: 6px;}
      .ok {background: rgba(31,122,77,.15); color: #1f7a4d;} .warn {background: rgba(154,91,0,.15); color: #9a5b00;}
    </style>
    """, unsafe_allow_html=True)

st.markdown('<div class="eyebrow">PF-3872 Metodologías Ágiles · UCR · Matriz multicriterio</div>', unsafe_allow_html=True)
st.title("Asesor de Agilidad a Escala")
st.caption("La IA lee el caso y llena la ficha → un motor de reglas calcula con la matriz del equipo (R1–R9) → "
           "el Excel descargable permite auditar cada número.")

hay_ia = _secretos_ok()
izq, der = st.columns([1.25, 1], gap="large")

with izq:
    st.subheader("Caso")
    opciones = list(CASOS) + ["Caso nuevo"]
    sel = st.radio("Casos del curso", opciones, horizontal=True, key="caso_sel", label_visibility="collapsed")
    if sel != st.session_state.get("caso_cargado"):
        st.session_state.caso_cargado = sel
        if sel == "Caso nuevo":
            cargar_en_estado("Caso nuevo", "", ficha_vacia(), "Ficha vacía")
        else:
            c = CASOS[sel]
            cargar_en_estado(c["titulo"], c["texto"], ficha_vacia(), "Ficha vacía")
    texto = st.text_area("Descripción del caso", key=f"texto_{st.session_state.version}",
                         value=st.session_state.texto, height=140,
                         placeholder="Número de equipos, productos, dependencias, regulación, portafolio, autonomía, tipos de trabajo…")
    st.session_state.texto = texto
    if hay_ia:
        if st.button("Analizar con IA", type="primary"):
            if len(texto.strip()) < 40:
                st.warning("Describa el caso con más detalle antes de analizar.")
            else:
                listo = False
                with st.spinner("La IA está leyendo el caso…"):
                    try:
                        analizar_con_ia(texto, "Caso evaluado" if sel == "Caso nuevo" else st.session_state.titulo)
                        listo = True
                    except Exception as e:  # noqa: BLE001
                        st.error(mensaje_error(e))
                if listo:
                    st.rerun()
    else:
        st.info("La lectura con IA no está configurada en este despliegue (falta GEMINI_API_KEY). "
                "Complete la ficha a mano: el cálculo y el Excel funcionan igual.")

    st.subheader("Ficha de captura")
    st.caption(f"Origen: **{st.session_state.origen}** · Explícito: el caso lo dice · Inferido: se deduce (anote el supuesto) · "
               "Faltante: sin datos, 2,5 para los cuatro marcos.")
    v = st.session_state.version
    for cid, c in reglas()["criterios"].items():
        d = st.session_state.ficha[cid]
        with st.container(border=True):
            st.markdown(f"**{cid} · {c['nombre']}** · peso {c['peso'] * 100:.0f} %" + (" · :red[excluyente]" if c["excluyente"] else ""))
            a, b = st.columns([1.4, 0.8])
            if cid == "C1":
                val = a.number_input("Número de equipos", min_value=0, step=1, key=f"v_{cid}_{v}",
                                     value=int(d["valor"]) if d["valor"] not in (None, "") else 0)
                d["valor"] = val or None
            else:
                ops = ["(sin nivel)"] + list(c["niveles"])
                actual = d["valor"] if d["valor"] in c["niveles"] else "(sin nivel)"
                val = a.selectbox("Nivel", ops, index=ops.index(actual), key=f"v_{cid}_{v}")
                d["valor"] = None if val == "(sin nivel)" else val
            est = b.selectbox("Estado", ESTADOS, index=ESTADOS.index(d["estado"]) if d["estado"] in ESTADOS else 2, key=f"e_{cid}_{v}")
            d["estado"] = est
            d["evidencia"] = st.text_input("Evidencia citada del caso", d["evidencia"], key=f"ev_{cid}_{v}")
            d["supuesto"] = st.text_input("Supuesto o justificación", d["supuesto"], key=f"su_{cid}_{v}")
            if d["estado"] != "Faltante" and d["valor"]:
                nivel = motor.banda_equipos(int(d["valor"])) if cid == "C1" else d["valor"]
                st.caption(f"Disparador · {nivel}: {DISP.get((cid, nivel), '')}")

    with st.expander("Restricciones del caso que descartan un marco (R4)"):
        for m in motor.MARCOS:
            on = st.checkbox(f"Excluir {m}", value=m in st.session_state.excl, key=f"x_{m}_{v}")
            motivo = st.text_input(f"Motivo ({m})", st.session_state.excl.get(m, ""), key=f"xm_{m}_{v}") if on else ""
            if on:
                st.session_state.excl[m] = motivo or "Restricción declarada en el caso"
            else:
                st.session_state.excl.pop(m, None)

# --------------------------------------------------------------------------- cálculo
ficha_sin_datos = all(d["estado"] == "Faltante" and d["valor"] in (None, "") for d in st.session_state.ficha.values())
if ficha_sin_datos:
    with der:
        st.markdown('<div class="veredicto" style="--c:#8C96AA"><div class="eyebrow">Sin evaluar</div>'
                    '<div class="marco">Esperando la ficha</div></div>', unsafe_allow_html=True)
        st.write("Pulse **Analizar con IA**: la IA leerá el caso, citará la evidencia de cada criterio y el motor "
                 "calculará el resultado. También puede llenar los siete criterios a mano." if hay_ia else
                 "Complete los siete criterios de la ficha a mano: el resultado se calcula al instante.")
    panel_bitacora()
    st.stop()
try:
    limpia = motor.normalizar_ficha(st.session_state.ficha, reglas())
    r = motor.calcular(limpia, reglas(), st.session_state.excl)
except ValueError as e:
    limpia, r = None, None
    with der:
        st.error(str(e))

if r:
    with der:
        if r["recomendacion"]:
            sensible = r["robustez"].startswith("SENSIBLE")
            pills = (f'<span class="pill {"warn" if sensible else "ok"}">{"Sensible: provisional" if sensible else "Robusta"}</span>'
                     + ('<span class="pill ok">Empate resuelto por R6</span>' if r["empate"] else ""))
            seg = (f" · 2.º {r['segundo']} {motor._fmt(r['totales'][r['segundo']])} · brecha {motor._fmt(r['brecha'])}"
                   if r["segundo"] else "")
            st.markdown(
                f'<div class="veredicto" style="--c:{COLORES[r["recomendacion"]]}"><div class="eyebrow">Recomendación'
                f'{" provisional" if sensible else ""}</div><div class="marco">{r["recomendacion"]}</div>'
                f'<div>{motor._fmt(r["totales"][r["recomendacion"]])} / 5,00{seg}</div><div style="margin-top:6px">{pills}</div></div>',
                unsafe_allow_html=True)
        else:
            st.error("Ninguno compatible (R9): los cuatro marcos quedaron excluidos.")
        antes = st.session_state.get("antes")
        if antes:
            ahora_ = instantanea(r)
            cambios = [f"{fa['id']} {fa['nivel']} → {fb['nivel']}" for fa, fb in zip(antes["filas"], ahora_["filas"])
                       if (fa["nivel"], fa["estado"]) != (fb["nivel"], fb["estado"])]
            st.success(f"**Tras su aclaración:** {_txt_total(antes)} ({_robustez_corta(antes).lower()}) → "
                       f"**{_txt_total(ahora_)}** ({_robustez_corta(ahora_).lower()}). "
                       + ("Criterios que cambiaron: " + "; ".join(cambios) + "." if cambios else "La ficha no cambió."))
        cfg = {"displayModeBar": False, "locale": "es"}
        t1, t2, t3 = st.tabs(["Puntaje total", "Desglose por criterio", "Perfil de compatibilidad"])
        with t1:
            st.plotly_chart(graficos.puntaje_total(r), width="stretch", config=cfg)
        with t2:
            st.caption("Cada segmento es peso × compatibilidad de un criterio. Pase el cursor para ver el cálculo; "
                       "toque un criterio en la leyenda para ocultarlo.")
            st.plotly_chart(graficos.desglose(r), width="stretch", config=cfg)
        with t3:
            st.caption("Compatibilidad 0–5 de cada marco en cada criterio. Toque un marco en la leyenda para compararlo.")
            st.plotly_chart(graficos.perfil(r), width="stretch", config=cfg)
        c = r["cobertura"]
        st.caption(f"Datos por peso: explícito {c['explícito']:.0%} · inferido {c['inferido']:.0%} · faltante {c['faltante']:.0%}")
        if r["recomendacion"] and r["segundo"]:
            st.markdown(f"**Qué separa a {r['recomendacion']} de {r['segundo']}**")
            st.markdown("\n".join(f"- {cid} {n}: **{'+' if d > 0 else ''}{motor._fmt(d)}**" for cid, n, d in r["aportes_vs_segundo"][:4]))
            if r["riesgos"]:
                st.caption("Riesgos de adopción (compatibilidad ≤ 2): " + "; ".join(f"{cid} {n}" for cid, n, _ in r["riesgos"]))

        st.subheader("Justificación")
        auto = st.session_state.pop("pedir_justificacion", False)
        if hay_ia and r["recomendacion"]:
            if auto or st.button("Redactar con IA"):
                cliente, modelo = gemini()
                pj = prompt_justificacion(st.session_state.texto, r)
                with st.spinner("Redactando…"):
                    try:
                        st.session_state.justificacion = cliente.models.generate_content(model=modelo, contents=pj).text
                    except Exception as e:  # noqa: BLE001
                        st.session_state.justificacion = f"No se pudo redactar la justificación: {mensaje_error(e)}"
                if st.session_state.get("ia"):
                    st.session_state.ia["prompt_justificacion"] = pj
                registrar(r, limpia)
        elif auto:
            registrar(r, limpia)
        st.write(st.session_state.justificacion or "_Use «Analizar con IA» o «Redactar con IA» para obtener la explicación._")

        if hay_ia and r.get("faltantes"):
            nombres = [f"{cid} {reglas()['criterios'][cid]['nombre']}" for cid in r["faltantes"]]
            st.subheader("Responder al agente")
            st.caption(("La recomendación es **provisional**. " if r["robustez"].startswith("SENSIBLE") else "")
                       + "Faltan datos de: " + "; ".join(nombres) + ". Si los conoce, respóndalos: se agregan al caso como "
                       "aclaración, la IA vuelve a leerlo y verá el resultado antes y después.")
            respuesta = st.text_area("Su respuesta a las preguntas del agente", key=f"acl_{v}", height=90,
                                     placeholder="Ej.: serán 20 equipos y la iniciativa formará parte del portafolio estratégico.")
            if st.button("Enviar respuesta y recalcular", disabled=not respuesta.strip()):
                previo, listo = instantanea(r), False
                with st.spinner("La IA está releyendo el caso con su aclaración…"):
                    try:
                        analizar_con_ia(st.session_state.texto.rstrip() + "\n\nAclaración del cliente: " + respuesta.strip(),
                                        st.session_state.titulo, "Ficha generada por IA tras la aclaración")
                        st.session_state.antes = previo
                        listo = True
                    except Exception as e:  # noqa: BLE001
                        st.error(mensaje_error(e))
                if listo:
                    st.rerun()

        with st.expander("Transparencia del agente: qué respondió la IA"):
            ia = st.session_state.get("ia")
            if not ia:
                st.write("Esta ficha no viene de la IA: se llenó a mano o todavía no se ha analizado el caso.")
            else:
                st.markdown(f"Modelo **{ia['modelo']}** · temperatura {ia['temperatura']} · {ia['hora']} (hora de Costa Rica).")
                st.caption("La app no guarda fichas ni resultados de los casos: cada análisis envía el enunciado a la IA y la ficha "
                           "de captura es esta respuesta, validada por el motor. Si después editó la ficha a mano, puede diferir.")
                for a in ia["ajustes"]:
                    st.warning(a)
                st.markdown("**Respuesta de la IA (JSON)**")
                try:
                    st.code(json.dumps(json.loads(ia["respuesta"]), ensure_ascii=False, indent=2), language="json")
                except ValueError:
                    st.code(ia["respuesta"])
                st.markdown("**Instrucciones enviadas para leer el caso**")
                st.code(ia["prompt"], language="markdown", wrap_lines=True)
                if ia.get("prompt_justificacion"):
                    st.markdown("**Instrucciones enviadas para redactar la justificación** (el resultado ya está calculado; la IA no puede cambiar números)")
                    st.code(ia["prompt_justificacion"], language="markdown", wrap_lines=True)

        st.subheader("Auditoría")
        with tempfile.TemporaryDirectory() as tmp:
            ruta = motor.exportar_excel(limpia, reglas(), os.path.join(tmp, "m.xlsx"), st.session_state.titulo,
                                        st.session_state.texto, st.session_state.excl)
            datos_xlsx = Path(ruta).read_bytes()
        base = motor._nombre_archivo(st.session_state.titulo)
        st.download_button("Descargar Excel para auditar", datos_xlsx, file_name=f"Matriz_{base}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
        st.download_button("Descargar ficha (JSON)", json.dumps(
            {"caso": st.session_state.titulo, "descripcion": st.session_state.texto, "origen": st.session_state.origen,
             "ficha": limpia, "exclusiones": st.session_state.excl, "totales": r["totales"],
             "recomendacion": r["recomendacion"], "robustez": r["robustez"]}, ensure_ascii=False, indent=2),
            file_name=f"Ficha_{base}.json", mime="application/json")
        st.caption("Abra la hoja Caso_Evaluado: Excel recalcula con sus fórmulas y debe mostrar los mismos totales.")
        b = st.session_state.bitacora
        if not (b and b[-1]["firma"] == firma(limpia)):
            if st.button("Registrar esta ejecución en la bitácora"):
                registrar(r, limpia)
                st.rerun()

    with izq:
        st.subheader("Cálculo trazable")
        filas = []
        for f in r["filas"]:
            fila = {"ID": f["id"], "Nivel": (f"{f['valor']} → {f['nivel']}" if f["id"] == "C1" and f["valor"] else f["nivel"]),
                    "Estado": f["estado"], "Peso": motor._fmt(f["peso"])}
            for m in motor.MARCOS:
                fila[motor.CORTO[m]] = f"{motor._num(f['puntos'][m])} → {motor._fmt(f['aportes'][m])}"
            filas.append(fila)
        filas.append({"ID": "Total", "Nivel": "", "Estado": "", "Peso": "1,00",
                      **{motor.CORTO[m]: motor._fmt(r["totales"][m]) for m in motor.MARCOS}})
        st.dataframe(pd.DataFrame(filas), hide_index=True, width="stretch")
        excl_txt = [f"**{m}** excluido (R3): compatibilidad 0 en {', '.join(r['auto_excl'][m])}" for m in motor.MARCOS if r["auto_excl"][m]]
        excl_txt += [f"**{m}** excluido (R4): {r['excl_manual'][m]}" for m in motor.MARCOS if m in r["excl_manual"]]
        st.markdown("**Restricciones críticas:** " + ("; ".join(excl_txt) if excl_txt else "ninguna, los cuatro marcos son elegibles."))
        st.markdown(f"**Robustez (R7):** {r['robustez']}")
        if r.get("faltantes"):
            st.dataframe(pd.DataFrame([{"Competidor": x, "Brecha actual": motor._fmt(v["brecha"]),
                                        "Ajuste peor caso": motor._fmt(v["ajuste"]), "Brecha peor caso": motor._fmt(v["peor"])}
                                       for x, v in r["peor_brecha"].items()]), hide_index=True, width="stretch")

with st.expander("Reglas de decisión de la matriz"):
    st.markdown(
        "- **R2 Cálculo:** Total = Σ peso × compatibilidad (máx. 5,00). Dato faltante = 2,5 para los cuatro marcos.\n"
        "- **R3 Exclusión automática:** compatibilidad 0 en un criterio excluyente (C1, C2, C7) con dato explícito o inferido.\n"
        "- **R4 Exclusión manual:** restricción declarada en el caso.\n"
        "- **R6 Empate técnico (≤ 0,10):** menos criterios con compatibilidad ≤ 1; luego mayor puntaje en C1+C2+C7; luego el marco más liviano.\n"
        "- **R7 Robustez:** brecha actual + Σ peso × menor diferencia posible en cada criterio faltante; si alguna queda ≤ 0, la recomendación es sensible.\n\n"
        "Pesos, niveles y puntajes se leen de `Matriz_Decision_Agilidad_Escala.xlsx`.")

panel_bitacora()
